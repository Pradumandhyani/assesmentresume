"""
Output infrastructure for the assignment generator.
Reads all per-topic JSON files and merges into final Excel + JSON output.
"""
import json, os, glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = r"d:\myProjects\redResumes\Assignment"
TOPICS_DIR = os.path.join(OUTPUT_DIR, "topics")

FIELDS = [
    "question_id", "title", "topic", "subtopic", "difficulty",
    "problem_statement", "input_format", "output_format", "constraints",
    "sample_input", "sample_output", "explanation",
    "test_cases", "edge_cases",
    "reference_solution_python", "reference_solution_java", "reference_solution_cpp",
    "time_complexity", "space_complexity"
]

TOPIC_COLORS = {
    "Array":                 "1E3A5F",
    "String":                "1B4332",
    "LinkedList":            "3D1A78",
    "Binary Trees":          "7B2D00",
    "Binary Search Trees":   "4A1942",
    "Stacks & Queues":       "1A3A4A",
    "Searching & Sorting":   "2D4A1A",
    "Greedy":                "4A3A1A",
    "BackTracking":          "4A1A1A",
    "Dynamic Programming":   "1A1A4A",
    "Graph":                 "1A4A3A",
    "Heap":                  "3A1A4A",
    "Trie":                  "1A4A1A",
    "Matrix":                "4A2A1A",
    "Bit Manipulation":      "1A3A3A",
}

def load_all_topics():
    all_problems = []
    for fp in sorted(glob.glob(os.path.join(TOPICS_DIR, "*.json"))):
        with open(fp, encoding="utf-8") as f:
            problems = json.load(f)
            all_problems.extend(problems)
    return all_problems

def serialize_field(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return str(v) if v is not None else ""

def write_json(problems):
    out = os.path.join(OUTPUT_DIR, "question_bank.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(problems, f, indent=2, ensure_ascii=False)
    print(f"JSON written: {out} ({len(problems)} problems)")

def write_excel(problems):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Question Bank"
    header_fill = PatternFill("solid", fgColor="0D1B2A")
    header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    thin = Side(style="thin", color="444444")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col_idx, field in enumerate(FIELDS, 1):
        cell = ws.cell(row=1, column=col_idx, value=field.upper().replace("_", " "))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28
    for row_idx, prob in enumerate(problems, 2):
        topic = prob.get("topic", "")
        color = TOPIC_COLORS.get(topic, "2A2A2A")
        row_fill = PatternFill("solid", fgColor=color)
        for col_idx, field in enumerate(FIELDS, 1):
            val = prob.get(field, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=serialize_field(val))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Calibri", size=10, color="F0F0F0")
            cell.fill = row_fill
            cell.border = border
        ws.row_dimensions[row_idx].height = 80
    col_widths = {1:14,2:35,3:22,4:25,5:12,6:70,7:40,8:40,9:40,10:40,11:30,12:50,13:60,14:50,15:60,16:60,17:60,18:22,19:22}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Summary")
    ws2["A1"]="Topic"; ws2["B1"]="Count"; ws2["C1"]="Easy"; ws2["D1"]="Medium"; ws2["E1"]="Hard"
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D1B2A")
        cell.alignment = Alignment(horizontal="center")
    topic_stats = {}
    for p in problems:
        t = p.get("topic","Unknown"); d = p.get("difficulty","Medium")
        if t not in topic_stats: topic_stats[t] = {"Easy":0,"Medium":0,"Hard":0}
        topic_stats[t][d] = topic_stats[t].get(d,0) + 1
    for r,(t,stats) in enumerate(sorted(topic_stats.items()),2):
        ws2.cell(r,1,t); ws2.cell(r,2,sum(stats.values()))
        ws2.cell(r,3,stats.get("Easy",0)); ws2.cell(r,4,stats.get("Medium",0)); ws2.cell(r,5,stats.get("Hard",0))
    out = os.path.join(OUTPUT_DIR, "question_bank.xlsx")
    wb.save(out)
    print(f"Excel written: {out} ({len(problems)} problems)")

if __name__ == "__main__":
    os.makedirs(TOPICS_DIR, exist_ok=True)
    problems = load_all_topics()
    if not problems:
        print("No topic JSON files found in", TOPICS_DIR)
    else:
        write_json(problems)
        write_excel(problems)
        print("Done.")
